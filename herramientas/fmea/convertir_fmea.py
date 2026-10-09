"""Convierte la memoria RCM (FMECA) y la lista de síntomas por condición a un borrador
de FMEA en el formato del modelo, más un Excel de revisión para Confiabilidad.

Uso (con las fuentes versionadas en knowledge/fmea/fuentes/):
    python herramientas/fmea/convertir_fmea.py

Todo lo generado queda en estado "propuesto": la tabla de cruce modo<->síntoma, la
severidad y el diccionario de variables requieren aprobación de Confiabilidad.
"""
import argparse
import csv
import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
MATRIZ_SMCS = RAIZ / "knowledge/smcs/Matriz_MF_SMCS_MAESTRO_CONSOLIDADO_v2.csv"

# Columnas de la hoja FMECA de la memoria RCM (fila 1-2 encabezados, datos desde fila 3)
COL = dict(sistema=1, subunidad=3, funcion=4, ff_id=5, falla_funcional=6, mf_id=7, modo=8,
           efecto=9, H=10, S=11, E=12, O=13, BC=14, RC=15, SC=16, OT=17, tarea=18,
           frecuencia=19, detencion=20, ejecutor=21)

# Diccionario síntoma -> (variable, dirección, unidad). Se aplica sobre "modo | síntoma".
# Elementos de laboratorio/particulado: se toman todas las coincidencias (un síntoma puede
# reportar varios elementos). Parámetros operacionales: solo la primera coincidencia.
DICCIONARIO_LAB = [
    (r"\blead\b|\bPb\b", "pb_ppm", "aumento", "ppm"),
    (r"sodium|sodim|\bNa\b", "na_ppm", "aumento", "ppm"),
    (r"potassium|potassiuk|\(K\)", "k_ppm", "aumento", "ppm"),
    (r"cop+er|cupper|\bCu\b", "cu_ppm", "aumento", "ppm"),
    (r"\biron\b|\bFe\b", "fe_ppm", "aumento", "ppm"),
    (r"silicon|silica|\bSi\b|dust", "si_ppm", "aumento", "ppm"),
    (r"soot|stoot", "hollin", "aumento", "%"),
    (r"\bPQI\b", "pqi", "aumento", "indice"),
    (r"particle count", "conteo_particulas", "aumento", "ISO 4406"),
    (r"filter cut|patch test with particulate|screen with|metallic particulate|particulate level",
     "particulado_inspeccion", "aumento", "nivel"),
    (r"viscosity", "viscosidad", "disminucion", "cSt"),
]
DICCIONARIO_OPER = [
    (r"\bDEF\b|aftertreatment", "presion_def", "disminucion", "kPa"),
    (r"crank\w*.*pressure", "presion_carter", "aumento", "kPa"),
    (r"oil.*level|level.*oil", "nivel_aceite", "disminucion", "nivel"),
    (r"oil leak", "fuga_aceite", "aumento", "evento"),
    (r"pressure differential", "presion_aceite_diferencial", "aumento", "kPa"),
    (r"oil.*pressure|lubrication pressure", "presion_aceite", "disminucion", "kPa"),
    (r"intake manifold.*pressure|booster pressure|IMAP", "presion_admision", "desviacion", "kPa"),
    (r"intake manifold.*temp|IMAT", "temp_admision", "aumento", "°C"),
    (r"exhaust.*differential|temp differential", "temp_escape_diferencial", "aumento", "°C"),
    (r"exhaust.*temp|temp.*exhaust", "temp_escape", "aumento", "°C"),
    (r"fuel.*temp", "temp_combustible", "aumento", "°C"),
    (r"coolant pressure", "presion_refrigerante", "disminucion", "kPa"),
    (r"coolant.*temperature low|low coolant temp|coolant low temp", "temp_refrigerante", "disminucion", "°C"),
    (r"coolant.*temp|overheat|over temperature", "temp_refrigerante", "aumento", "°C"),
    (r"coolant.*level|level.*coolant|coolant leak", "nivel_refrigerante", "disminucion", "nivel"),
    (r"oil.*temp", "temp_aceite", "aumento", "°C"),
    (r"oil filter", "restriccion_filtro_aceite", "aumento", "kPa"),
    (r"air filter", "restriccion_filtro_aire", "aumento", "kPa"),
    (r"fan (drive|motor).*filter", "restriccion_filtro_ventilador", "aumento", "kPa"),
    (r"h2o|water", "agua_en_combustible", "aumento", "nivel"),
    (r"fuel.*(filter|restrict)", "restriccion_filtro_combustible", "aumento", "kPa"),
    (r"fuel.*pressure|\bFCV\b", "presion_combustible", "desviacion", "kPa"),
    (r"fan.*speed|overspeed", "velocidad_ventilador", "desviacion", "rpm"),
]
INSTRUMENTACION = re.compile(
    r"ECM|communication|wrong data|voltage|erratic|sensor|^electric|current (low|high)", re.I)
CODIGO_DIAGNOSTICO = re.compile(r"FMI", re.I)

# Modos predictivos propuestos (nivel intermedio entre la memoria RCM y los síntomas).
# Padres RCM: (código SMCS de la subunidad, regex sobre el texto del modo RCM).
# Las señales salen de los planes de acción del catálogo de síntomas; requieren validación.
MODOS_PREDICTIVOS = [
    ("MFP-01", "desgaste_cojinetes", [("1000", "DESGASTE"), ("1319", "DESGASTE")],
     [("pb_ppm", "aumento"), ("cu_ppm", "aumento"), ("pqi", "aumento"), ("particulado_inspeccion", "aumento"),
      ("presion_aceite", "disminucion")],
     "Plan Cat: 'Elevated lead alone is often related to sleeve bearings'."),
    ("MFP-02", "desgaste_camisas_anillos", [("1000", "DESGASTE")],
     [("fe_ppm", "aumento"), ("presion_carter", "aumento"), ("hollin", "aumento"), ("pqi", "aumento"),
      ("particulado_inspeccion", "aumento")],
     "Plan Cat: 'Iron can be sourced to Liners, crankshaft, camshaft'; presión de cárter alta = blowby."),
    ("MFP-03", "contaminacion_refrigerante_en_aceite", [("1000", "DESGASTE")],
     [("na_ppm", "aumento"), ("k_ppm", "aumento"), ("cu_ppm", "aumento"), ("nivel_refrigerante", "disminucion")],
     "Plan Cat: 'Elevated sodium is often related to coolant contamination'; síntomas 'Coolant Transfer'."),
    ("MFP-04", "ingreso_abrasivo_polvo", [("1054", "SATURACION"), ("1050", "RESTRICCION")],
     [("si_ppm", "aumento"), ("restriccion_filtro_aire", "aumento"), ("fe_ppm", "aumento")],
     "Síntomas 'Silicon high' y 'Contamination by Dust'."),
    ("MFP-05", "dilucion_combustible", [("1290", "DESGASTE"), ("1251", "DESGASTE")],
     [("viscosidad", "disminucion")],
     "Plan Cat: 'Fuel contamination can cause lower viscosity'."),
    ("MFP-06", "sobrecalentamiento", [("1000", "DESGASTE")],
     [("temp_refrigerante", "aumento"), ("temp_aceite", "aumento"), ("nivel_refrigerante", "disminucion"),
      ("presion_refrigerante", "disminucion")],
     "Síntomas VIMS de temperatura de refrigerante/aceite y nivel de refrigerante."),
    ("MFP-07", "baja_presion_lubricacion",
     [("1000", "NIVEL DE ACEITE|FUGAS DE ACEITE|FILTRO CENTRIFUGO"), ("1319", "DESGASTE")],
     [("presion_aceite", "disminucion"), ("nivel_aceite", "disminucion"), ("viscosidad", "disminucion")],
     "Síntomas 'Engine Oil Pressure low' / 'Low oil pressure'. Usa presión de ACEITE, no de cárter."),
    ("MFP-08", "combustion_deficiente",
     [("1290", "DESGASTE"), ("1052", "DESGASTE"), ("1064", "DESGASTE"), ("1054", "SATURACION")],
     [("temp_escape", "aumento"), ("temp_escape_diferencial", "aumento"), ("hollin", "aumento"),
      ("temp_admision", "aumento"), ("presion_admision", "desviacion")],
     "Síntomas de escape, admisión (IMAT/boost) y hollín."),
]


def _txt(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else None


def _sn(v):
    v = _txt(v)
    return v.upper() if v else None


def severidad_sugerida(H, S, E, O):
    """Regla propuesta desde la hoja de decisión RCM (a validar por Confiabilidad)."""
    if S == "S" or E == "S":
        return "critica", "consecuencia de seguridad o medio ambiente"
    if O == "S":
        return "alta", "consecuencia operacional"
    if H == "N":
        return "media", "falla oculta (riesgo de falla múltiple)"
    return "baja", "evidente sin consecuencia operacional"


def intervalo_horas(frecuencia):
    horas = re.findall(r"([\d.]+)\s*H\b", frecuencia or "", re.I)
    return int(horas[-1].replace(".", "")) if horas else None


def cargar_matriz_smcs():
    codigos = defaultdict(set)
    if MATRIZ_SMCS.exists():
        with open(MATRIZ_SMCS, encoding="utf-8-sig") as f:
            for fila in csv.DictReader(f):
                codigos[fila["Codigo_SMCS"].strip()].add(
                    (fila["Sistema"].strip(), fila["Descripcion_SMCS"].strip()))
    return codigos


def leer_memoria(ruta, prefijo_sistema):
    ws = openpyxl.load_workbook(ruta, data_only=True, read_only=True)["FMECA"]
    smcs = cargar_matriz_smcs()
    modos, ctx = [], {}
    for n, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if n < 3:
            continue
        r = list(r) + [None] * (22 - len(r))
        for k in ("sistema", "subunidad", "funcion", "falla_funcional"):
            if _txt(r[COL[k]]):
                ctx[k] = _txt(r[COL[k]])
                if k == "sistema":
                    ctx.pop("subunidad", None)
        if not _txt(r[COL["modo"]]) or not re.search(prefijo_sistema, ctx.get("sistema") or ""):
            continue
        H, S, E, O = (_sn(r[COL[c]]) for c in "HSEO")
        sev, regla = severidad_sugerida(H, S, E, O)
        m = re.search(r"\b(\d{3}[0-9A-Z])\b", ctx.get("subunidad", ""))
        codigo = m.group(1) if m else None
        en_matriz = sorted(smcs.get(codigo, []))
        modos.append({
            "id": f"RCM-{len(modos) + 1:03d}",
            "sistema": ctx.get("sistema"), "subunidad": ctx.get("subunidad"),
            "codigo_smcs": codigo,
            "smcs_en_matriz": "; ".join(f"{s}: {d}" for s, d in en_matriz) or "NO ENCONTRADO",
            "funcion": ctx.get("funcion"), "falla_funcional": ctx.get("falla_funcional"),
            "modo_falla": _txt(r[COL["modo"]]), "efecto": _txt(r[COL["efecto"]]),
            "evidente": {"S": True, "N": False}.get(H),
            "consecuencia_rcm": "".join(c for c, v in zip("SEO", (S, E, O)) if v == "S") or "-",
            "severidad_sugerida": sev, "regla_severidad": regla,
            "tipo_tarea": next((t for t in ("BC", "RC", "SC", "OT") if _sn(r[COL[t]]) == "S"), None),
            "tarea": _txt(r[COL["tarea"]]), "frecuencia": _txt(r[COL["frecuencia"]]),
            "intervalo_h": intervalo_horas(_txt(r[COL["frecuencia"]])),
            "origen": {"archivo": Path(ruta).name, "hoja": "FMECA", "fila": n},
            "estado_revision": "propuesto",
        })
    return modos


def mapear_sintoma(modo, sintoma):
    """Devuelve la lista de señales [(variable, dirección, unidad)] de un síntoma."""
    texto = f"{modo or ''} | {sintoma or ''}"
    lab = [(v, d, u) for patron, v, d, u in DICCIONARIO_LAB if re.search(patron, texto, re.I)]
    if lab:
        return lab
    for patron, v, d, u in DICCIONARIO_OPER:
        if re.search(patron, texto, re.I):
            return [(v, d, u)]
    return []


def leer_sintomas(ruta, filtro_componente):
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    sintomas = {}
    for hoja in wb.sheetnames:
        filas = list(wb[hoja].iter_rows(values_only=True))
        enc = next((i for i, f in enumerate(filas[:10]) if "MODELS" in f and "SYMPTOM" in f), None)
        if enc is None:
            continue
        for ini in [i for i, v in enumerate(filas[enc]) if v == "MODELS"]:
            for n, f in enumerate(filas[enc + 1:], start=enc + 2):
                f = list(f) + [None] * 6
                modelos, comp, rutina, modo, sintoma, plan = (_txt(x) for x in f[ini:ini + 6])
                if not comp or not re.search(filtro_componente, comp, re.I):
                    continue
                clave = (rutina, (modo or "").lower(), (sintoma or "").lower())
                origen = {"archivo": Path(ruta).name, "hoja": hoja, "fila": n}
                mods = {m.strip() for m in (modelos or "").split(",") if m.strip()}
                if clave in sintomas:
                    sintomas[clave]["modelos"] |= mods
                    sintomas[clave]["origen"].append(origen)
                    continue
                texto = f"{modo or ''} {sintoma or ''}"
                if INSTRUMENTACION.search(texto):
                    categoria, senales = "instrumentacion", []
                elif CODIGO_DIAGNOSTICO.search(texto):
                    categoria, senales = "evento_diagnostico", []
                else:
                    senales = mapear_sintoma(modo, sintoma)
                    categoria = {"SOS": "aceite_lab", "Particulate": "particulado"}.get(
                        rutina, "parametro_operacional")
                sintomas[clave] = {
                    "modelos": mods, "componente": comp, "rutina": rutina,
                    "modo_origen": modo, "sintoma": sintoma, "plan_accion": plan,
                    "categoria": categoria,
                    "senales": [{"variable": v, "direccion": d, "unidad": u} for v, d, u in senales],
                    "origen": [origen],
                }
    lista = []
    for i, s in enumerate(sorted(sintomas.values(), key=lambda s: (s["categoria"], s["rutina"] or "")), 1):
        s["id"] = f"SIN-{i:03d}"
        s["modelos"] = sorted(s["modelos"])
        propias = {(x["variable"], x["direccion"]) for x in s["senales"]}
        s["modos_predictivos"] = [mid for mid, _, _, senales, _ in MODOS_PREDICTIVOS
                                  if propias & set(senales)]
        lista.append(s)
    return lista


def construir_modos_predictivos(modos_rcm, sintomas):
    def padres_de(reglas):
        return [m for codigo, patron in reglas for m in modos_rcm
                if m["codigo_smcs"] == codigo and re.search(patron, m["modo_falla"] or "", re.I)]

    orden = ["baja", "media", "alta", "critica"]
    salida = []
    for mid, nombre, reglas, senales, evidencia in MODOS_PREDICTIVOS:
        padres = padres_de(reglas)
        codigos = sorted({c for c, _ in reglas})
        sev = max((m["severidad_sugerida"] for m in padres), key=orden.index, default="alta")
        soporte = [s for s in sintomas if mid in s["modos_predictivos"]]
        salida.append({
            "id": mid, "nombre": nombre, "codigos_smcs": codigos,
            "modos_rcm_padre": [m["id"] for m in padres],
            "severidad": sev,
            "variables": sorted({v for v, _ in senales}),
            "senales": [{"variable": v, "direccion": d} for v, d in senales],
            "precursores": sorted({f"{s['sintoma']} ({s['rutina']})" for s in soporte}),
            "sintomas_soporte": [s["id"] for s in soporte],
            "modelos_aplicables": sorted({m for s in soporte for m in s["modelos"]}),
            "evidencia": evidencia,
            "estado_revision": "propuesto",
        })
    return salida


def escribir_excel(ruta, modos_rcm, sintomas, predictivos):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    fuente, negrita = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    amarillo = PatternFill("solid", start_color="FFFF00")
    gris = PatternFill("solid", start_color="D9D9D9")
    wb = openpyxl.Workbook()

    ins = wb.active
    ins.title = "Instrucciones"
    textos = [
        ("Revisión FMEA para el modelo predictivo – Motor CAEX", negrita),
        ("Todo el contenido está en estado PROPUESTO y fue generado automáticamente desde la "
         "memoria RCM y la lista de síntomas por condición.", fuente),
        ("Qué revisar:", negrita),
        ("1. 'Modos predictivos': modos que usará el modelo, sus señales (variable + dirección) y severidad.", fuente),
        ("2. 'Síntomas': variable y dirección asignadas a cada síntoma y a qué modo predictivo aportan.", fuente),
        ("3. 'Modos RCM': severidad sugerida y código SMCS de cada modo de la memoria.", fuente),
        ("Cómo responder: completar solo las columnas con fondo AMARILLO.", negrita),
        ("  Aprobado: S / N.   Corrección: valor correcto si corresponde.   Comentario: libre.", fuente),
        ("  Ejemplo: Aprobado = N | Corrección = 'pb_ppm, aumento' | Comentario = 'el plomo indica cojinetes, no bujes'", fuente),
        ("Regla de severidad sugerida (a validar):", negrita),
        ("  S o E = 'S' → critica · O = 'S' → alta · falla oculta (H = 'N') → media · resto → baja", fuente),
        ("Severidad de modos predictivos: la mayor entre sus modos RCM padre (por código SMCS).", fuente),
    ]
    for i, (t, f) in enumerate(textos, 1):
        ins.cell(i, 1, t).font = f
    ins.column_dimensions["A"].width = 120

    def hoja(titulo, columnas, filas):
        ws = wb.create_sheet(titulo)
        cols = columnas + ["Aprobado", "Corrección", "Comentario"]
        for j, c in enumerate(cols, 1):
            celda = ws.cell(1, j, c)
            celda.font, celda.fill = negrita, (amarillo if j > len(columnas) else gris)
        for i, fila in enumerate(filas, 2):
            for j, v in enumerate(fila, 1):
                if isinstance(v, (list, dict)):
                    v = "; ".join(map(str, v)) if isinstance(v, list) else json.dumps(v, ensure_ascii=False)
                c = ws.cell(i, j, v)
                c.font, c.alignment = fuente, Alignment(wrap_text=True, vertical="top")
            for j in range(len(columnas) + 1, len(cols) + 1):
                c = ws.cell(i, j)
                c.fill, c.font = amarillo, fuente
        dv = DataValidation(type="list", formula1='"S,N"', allow_blank=True)
        ws.add_data_validation(dv)
        col = openpyxl.utils.get_column_letter(len(columnas) + 1)
        dv.add(f"{col}2:{col}{len(filas) + 1}")
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = min(
                60, max(10, len(c) + 2, *(len(str(ws.cell(i, j).value or ""))
                                          for i in range(2, min(len(filas) + 2, 40))))) * 0.9
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    hoja("Modos predictivos",
         ["ID", "Nombre", "Severidad", "Señales (variable:dirección)", "Códigos SMCS", "Modos RCM padre",
          "Síntomas soporte", "Modelos con evidencia", "Evidencia"],
         [[p["id"], p["nombre"], p["severidad"],
           [f"{s['variable']}:{s['direccion']}" for s in p["senales"]], p["codigos_smcs"],
           p["modos_rcm_padre"], p["sintomas_soporte"], p["modelos_aplicables"], p["evidencia"]]
          for p in predictivos])
    hoja("Síntomas",
         ["ID", "Categoría", "Rutina", "Modo (origen)", "Síntoma", "Variable", "Dirección", "Unidad",
          "Modos predictivos", "Modelos", "Plan de acción", "Origen"],
         [[s["id"], s["categoria"], s["rutina"], s["modo_origen"], s["sintoma"],
           [x["variable"] for x in s["senales"]], [x["direccion"] for x in s["senales"]],
           [x["unidad"] for x in s["senales"]], s["modos_predictivos"], s["modelos"], s["plan_accion"],
           [f"{o['hoja']}!{o['fila']}" for o in s["origen"]]] for s in sintomas])
    hoja("Modos RCM",
         ["ID", "Subunidad", "Código SMCS", "SMCS en matriz", "Falla funcional", "Modo de falla", "Efecto",
          "Evidente", "Consecuencia", "Severidad sugerida", "Regla", "Tipo tarea", "Tarea", "Intervalo (h)",
          "Origen"],
         [[m["id"], m["subunidad"], m["codigo_smcs"], m["smcs_en_matriz"], m["falla_funcional"],
           m["modo_falla"], m["efecto"], {True: "S", False: "N"}.get(m["evidente"]), m["consecuencia_rcm"],
           m["severidad_sugerida"], m["regla_severidad"], m["tipo_tarea"], m["tarea"], m["intervalo_h"],
           f"{m['origen']['hoja']}!{m['origen']['fila']}"] for m in modos_rcm])

    res = wb.create_sheet("Resumen", 1)
    res["A1"], res["B1"] = "Indicador", "Valor"
    filas_res = [
        ("Modos RCM del sistema", "=COUNTA('Modos RCM'!A:A)-1"),
        ("Modos RCM con SMCS no encontrado en la matriz", "=COUNTIF('Modos RCM'!D:D,\"NO ENCONTRADO\")"),
        ("Síntomas (sin duplicados)", "=COUNTA('Síntomas'!A:A)-1"),
        ("Síntomas de instrumentación (van a Calidad, no al FMEA)", "=COUNTIF('Síntomas'!B:B,\"instrumentacion\")"),
        ("Síntomas de código de diagnóstico FMI (revisar caso a caso)", "=COUNTIF('Síntomas'!B:B,\"evento_diagnostico\")"),
        ("Síntomas físicos sin variable asignada", "=COUNTIFS('Síntomas'!F2:F1000,\"\",'Síntomas'!B2:B1000,\"aceite_lab\")+COUNTIFS('Síntomas'!F2:F1000,\"\",'Síntomas'!B2:B1000,\"particulado\")+COUNTIFS('Síntomas'!F2:F1000,\"\",'Síntomas'!B2:B1000,\"parametro_operacional\")"),
        ("Modos predictivos propuestos", "=COUNTA('Modos predictivos'!A:A)-1"),
        ("Ítems revisados (Aprobado S/N) en las tres hojas",
         "=COUNTA('Modos predictivos'!J:J)+COUNTA('Síntomas'!M:M)+COUNTA('Modos RCM'!P:P)-3"),
    ]
    for i, (k, f) in enumerate(filas_res, 2):
        res.cell(i, 1, k).font, res.cell(i, 2, f).font = fuente, fuente
    res["A1"].font = res["B1"].font = negrita
    res.column_dimensions["A"].width = 60
    wb.save(ruta)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--memoria", default=str(RAIZ / "knowledge/fmea/fuentes/FMEA_Memoria_798AC.xlsx"))
    ap.add_argument("--sintomas", default=str(RAIZ / "knowledge/fmea/fuentes/FMECA_Lists_ENG.xlsx"))
    ap.add_argument("--salida", default=str(RAIZ / "knowledge/fmea/borrador"))
    ap.add_argument("--sistema", default=r"\bENG\b|\bCS\b",
                    help="regex del sistema en la memoria RCM (motor y enfriamiento)")
    ap.add_argument("--componente", default=r"engine|motor", help="regex de componente en la lista de síntomas")
    a = ap.parse_args()

    modos_rcm = leer_memoria(a.memoria, a.sistema)
    sintomas = leer_sintomas(a.sintomas, a.componente)
    predictivos = construir_modos_predictivos(modos_rcm, sintomas)

    out = Path(a.salida)
    out.mkdir(parents=True, exist_ok=True)
    borrador = {
        "version": "1.0.0-borrador",
        "estado": "propuesto",
        "fuentes": [Path(a.memoria).name, Path(a.sintomas).name],
        "modos_falla": [{k: p[k] for k in ("id", "nombre", "severidad", "variables", "senales",
                                           "precursores", "codigos_smcs", "modos_rcm_padre", "sintomas_soporte",
                                           "modelos_aplicables", "estado_revision")}
                        for p in predictivos],
        "modos_rcm": modos_rcm,
        "sintomas": sintomas,
    }
    (out / "fmea_borrador.json").write_text(json.dumps(borrador, ensure_ascii=False, indent=2), "utf-8")
    escribir_excel(out / "revision_fmea_confiabilidad.xlsx", modos_rcm, sintomas, predictivos)
    print(f"modos RCM: {len(modos_rcm)} | síntomas: {len(sintomas)} | modos predictivos: {len(predictivos)}")
    print(f"salidas en {out}/")


if __name__ == "__main__":
    main()
