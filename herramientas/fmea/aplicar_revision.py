"""Aplica la planilla revisada por Confiabilidad y genera el FMEA aprobado.

Uso:
    python herramientas/fmea/aplicar_revision.py \
        --revision knowledge/fmea/revision/revision_fmea_confiabilidad_v1.xlsx

Reglas:
- Modos predictivos: se toma la versión de Confiabilidad (bloque encabezado por "ID:" si
  existe; si no, las filas originales). Solo entran los aprobados (S).
- Síntomas: S = aprobado, N = eliminado, otro valor (p. ej. P) = pendiente. El soporte de
  cada modo se recalcula solo con síntomas aprobados.
- Lo que la revisión afirma sin respaldo documental (modelos aplicables, códigos SMCS fuera
  de la matriz) se conserva, pero se marca como observación.
"""
import argparse
import csv
import json
import re
from datetime import date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
MATRIZ_SMCS = RAIZ / "knowledge/smcs/Matriz_MF_SMCS_MAESTRO_CONSOLIDADO_v2.csv"


def _lista(v):
    return [x.strip() for x in re.split(r"[;,]", str(v or "")) if x.strip()]


def _filas(ws):
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            yield dict(zip(hdr, r))


def leer_modos(ws):
    filas = list(ws.iter_rows(values_only=True))
    inicio = next((i for i, f in enumerate(filas) if f[0] == "ID:"), None)
    hdr = [c.value for c in ws[1]]
    bloque = filas[inicio + 1:] if inicio is not None else filas[1:]
    return [dict(zip(hdr, f)) for f in bloque if f[0] and str(f[0]).startswith("MFP-")], inicio is not None


def leer_sintomas(ws):
    sintomas = {}
    for f in _filas(ws):
        aprob = str(f.get("Aprobado") or "").strip().upper()
        estado = {"S": "aprobado", "N": "eliminado"}.get(aprob, "pendiente")
        vars_, dirs = _lista(f["Variable"]), _lista(f["Dirección"])
        sintomas[f["ID"]] = {
            "id": f["ID"], "estado": estado, "categoria": f["Categoría"], "rutina": f["Rutina"],
            "modo_origen": f["Modo (origen)"], "sintoma": f["Síntoma"],
            "senales": [{"variable": v, "direccion": d} for v, d in zip(vars_, dirs)],
            "modos_predictivos": _lista(f["Modos predictivos"]), "modelos": _lista(f["Modelos"]),
            "origen": _lista(f["Origen"]),
            "correccion": f.get("Corrección"), "comentario": f.get("Comentario"),
        }
    return sintomas


def codigos_matriz():
    with open(MATRIZ_SMCS, encoding="utf-8-sig") as f:
        return {fila["Codigo_SMCS"].strip() for fila in csv.DictReader(f)}


def construir(modos, sintomas, revision):
    en_matriz = codigos_matriz()
    aprobados = {k: s for k, s in sintomas.items() if s["estado"] == "aprobado"}
    salida, observaciones = [], []
    for m in modos:
        if str(m.get("Aprobado") or "").strip().upper() != "S":
            observaciones.append(f"{m['ID']}: no aprobado, excluido.")
            continue
        senales = [dict(zip(("variable", "direccion"), x.split(":")))
                   for x in _lista(m["Señales (variable:dirección)"])]
        soporte = [s for s in aprobados.values() if m["ID"] in s["modos_predictivos"]]
        declarado = set(re.findall(r"SIN-\d{3}", str(m["Síntomas soporte"])))
        quitados = sorted(declarado - {s["id"] for s in soporte})
        codigos = sorted(set(re.findall(r"\b\d{4}\b", str(m["Códigos SMCS"]))))
        con_evidencia = sorted({mod for s in soporte for mod in s["modelos"]})
        aplicables = _lista(m["Modelos con evidencia"])
        obs = []
        if quitados:
            obs.append("Síntomas de soporte declarados que no están aprobados: " + ", ".join(quitados))
        sin_evidencia = sorted(set(aplicables) - set(con_evidencia))
        if sin_evidencia:
            obs.append("Modelos aplicables por criterio experto, sin síntoma documental: "
                       + ", ".join(sin_evidencia))
        fuera = [c for c in codigos if c not in en_matriz]
        if fuera:
            obs.append("Códigos SMCS no presentes en la matriz: " + ", ".join(fuera))
        if "pendiente" in str(m["Modos RCM padre"]).lower():
            obs.append("Modos RCM padre pendientes de validar.")
        salida.append({
            "id": m["ID"], "nombre": m["Nombre"], "severidad": str(m["Severidad"]).strip().lower(),
            "variables": sorted({s["variable"] for s in senales}),
            "senales": senales,
            "precursores": sorted({f"{s['sintoma']} ({s['rutina']})" for s in soporte}),
            "codigos_smcs": codigos,
            "modos_rcm_padre": re.findall(r"RCM-\d{3}", str(m["Modos RCM padre"])),
            "sintomas_soporte": sorted(s["id"] for s in soporte),
            "modelos_aplicables": aplicables,
            "modelos_con_evidencia_documental": con_evidencia,
            "evidencia": m["Evidencia"],
            "observaciones": obs,
        })
    pendientes = [s for s in sintomas.values() if s["estado"] == "pendiente"]
    eliminados = [s for s in sintomas.values() if s["estado"] == "eliminado"]
    return {
        "version": "1.0.0",
        "componente": "motor_diesel",
        "estado": "aprobado",
        "aprobacion": {"revision": Path(revision).name, "fecha_aplicacion": date.today().isoformat(),
                       "aprobado_por": "Confiabilidad"},
        "modos_falla": salida,
        "sintomas_aprobados": sorted(aprobados.values(), key=lambda s: s["id"]),
        "sintomas_pendientes": sorted(pendientes, key=lambda s: s["id"]),
        "sintomas_eliminados": sorted(eliminados, key=lambda s: s["id"]),
        "observaciones": observaciones,
    }


def reporte(aprobado, revision_rcm_vacia, bloque_revisor):
    lineas = [f"# FMEA motor – revisión aplicada (v{aprobado['version']})", "",
              f"Revisión: `{aprobado['aprobacion']['revision']}`, aplicada el "
              f"{aprobado['aprobacion']['fecha_aplicacion']}.", "",
              "## Resultado", "",
              f"- Modos predictivos aprobados: {len(aprobado['modos_falla'])}",
              f"- Síntomas aprobados: {len(aprobado['sintomas_aprobados'])}",
              f"- Síntomas eliminados: {len(aprobado['sintomas_eliminados'])}",
              f"- Síntomas pendientes: {len(aprobado['sintomas_pendientes'])}", ""]
    if bloque_revisor:
        lineas += ["Los modos se tomaron del bloque editado por Confiabilidad (encabezado `ID:`).", ""]
    lineas += ["## Observaciones por modo", ""]
    for m in aprobado["modos_falla"]:
        lineas.append(f"### {m['id']} {m['nombre']} ({m['severidad']})")
        lineas += [f"- {o}" for o in m["observaciones"]] or ["- Sin observaciones."]
        lineas.append("")
    lineas += ["## Síntomas pendientes", "", "| ID | Síntoma | Corrección pedida | Comentario |",
               "|---|---|---|---|"]
    lineas += [f"| {s['id']} | {s['sintoma']} | {s['correccion'] or ''} | {s['comentario'] or ''} |"
               for s in aprobado["sintomas_pendientes"]]
    lineas += ["", "## Síntomas eliminados", "", "| ID | Síntoma | Motivo |", "|---|---|---|"]
    lineas += [f"| {s['id']} | {s['sintoma']} | {s['comentario'] or ''} |"
               for s in aprobado["sintomas_eliminados"]]
    if revision_rcm_vacia:
        lineas += ["", "## Modos RCM", "",
                   "La hoja 'Modos RCM' no fue revisada: la regla de severidad H/S/E/O sigue sin validar. "
                   "Las severidades de los modos predictivos son las asignadas por Confiabilidad."]
    return "\n".join(lineas) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--revision", required=True)
    ap.add_argument("--salida", default=str(RAIZ / "knowledge/fmea/aprobado"))
    ap.add_argument("--runtime", default=str(RAIZ / "components/motor_diesel/fmea.json"))
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.revision, data_only=True)
    modos, bloque = leer_modos(wb["Modos predictivos"])
    sintomas = leer_sintomas(wb["Síntomas"])
    rcm_vacia = not any(f.get("Aprobado") for f in _filas(wb["Modos RCM"]))
    aprobado = construir(modos, sintomas, a.revision)

    out = Path(a.salida)
    out.mkdir(parents=True, exist_ok=True)
    (out / "fmea_motor_v1.json").write_text(json.dumps(aprobado, ensure_ascii=False, indent=2), "utf-8")
    (out / "REPORTE_REVISION_v1.md").write_text(reporte(aprobado, rcm_vacia, bloque), "utf-8")

    runtime = {
        "version": aprobado["version"],
        "fuente": "knowledge/fmea/aprobado/fmea_motor_v1.json",
        "modos_falla": [{k: m[k] for k in ("id", "nombre", "severidad", "variables", "senales",
                                           "precursores", "codigos_smcs")}
                        for m in aprobado["modos_falla"]],
    }
    Path(a.runtime).write_text(json.dumps(runtime, ensure_ascii=False, indent=1) + "\n", "utf-8")
    print(f"modos aprobados: {len(aprobado['modos_falla'])} | síntomas aprobados: "
          f"{len(aprobado['sintomas_aprobados'])} | pendientes: {len(aprobado['sintomas_pendientes'])} | "
          f"eliminados: {len(aprobado['sintomas_eliminados'])}")


if __name__ == "__main__":
    main()
